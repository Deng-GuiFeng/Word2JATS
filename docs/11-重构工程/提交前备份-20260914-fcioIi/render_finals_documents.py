"""由技术说明 Markdown 生成 Word，并用同批验证结果生成答辩 PPT。"""
from pathlib import Path
import argparse
import json
import re
import statistics

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "决赛提交"


def plain(text):
    text = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1（\2）", text)
    return text.replace("**", "").replace("`", "")


def document(source=None):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(2.3)
    section.header_distance = Cm(.65)
    section.footer_distance = Cm(.8)
    section.left_margin, section.right_margin = Cm(2.0), Cm(2.0)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption", "List Bullet"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.element.rPr.rFonts.set(qn("w:eastAsia"),
                                    "Noto Sans CJK SC" if name in ("Title", "Heading 1", "Heading 2", "Heading 3")
                                    else "Noto Serif CJK SC")
        for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            style.element.rPr.rFonts.attrib.pop(qn('w:' + attr), None)
        language = OxmlElement('w:lang')
        language.set(qn('w:eastAsia'), 'zh-CN')
        style.element.rPr.append(language)
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.22
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.keep_together = False
    normal.paragraph_format.widow_control = True
    for name, size in (("Title", 25), ("Heading 1", 17), ("Heading 2", 12.5)):
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles['Caption'].font.size = Pt(9)
    doc.styles['Caption'].font.italic = False
    for border in doc.styles['Title'].element.findall('.//' + qn('w:bottom')):
        border.set(qn('w:color'), '808080')
        border.attrib.pop(qn('w:themeColor'), None)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    source = source or ROOT / "决赛提交/技术方案说明书.md"
    lines = source.read_text().splitlines()
    targets = {}
    for line in lines:
        if line.startswith('## '):
            targets[line[3:]] = f'section_{len(targets)}'
        match = re.match(r'^(?:!\[)?([图表]\d+)[　\s]', line)
        if match:
            label = match.group(1)
            targets[label] = ('figure_' if label.startswith('图') else 'table_') + label[1:]
    bookmark_id = 0

    def bookmark(paragraph, name):
        nonlocal bookmark_id
        start, end = OxmlElement('w:bookmarkStart'), OxmlElement('w:bookmarkEnd')
        bookmark_id += 1
        start.set(qn('w:id'), str(bookmark_id))
        start.set(qn('w:name'), name)
        end.set(qn('w:id'), str(bookmark_id))
        paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
        paragraph._p.append(end)

    def hyperlink(paragraph, label, target):
        element = OxmlElement('w:hyperlink')
        if target.startswith('#'):
            element.set(qn('w:anchor'), target[1:])
        else:
            element.set(qn('r:id'), paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True))
        element.set(qn('w:history'), '1')
        run, props = OxmlElement('w:r'), OxmlElement('w:rPr')
        fonts = OxmlElement('w:rFonts')
        for key, value in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia','Noto Serif CJK SC')]:
            fonts.set(qn('w:' + key), value)
        props.append(fonts)
        underline = OxmlElement('w:u')
        underline.set(qn('w:val'), 'single')
        props.append(underline)
        color = OxmlElement('w:color')
        color.set(qn('w:val'), '000000')
        props.append(color)
        run.append(props)
        text = OxmlElement('w:t')
        text.text = label
        run.append(text)
        element.append(run)
        paragraph._p.append(element)

    inline_pattern = re.compile(r'\[([^]]+)\]\(([^)]+)\)|https?://[^\s，。；）]+|\*\*([^*]+)\*\*|`([^`]+)`|[图表]\d+')

    def inline(paragraph, text):
        last = 0
        for match in inline_pattern.finditer(text):
            paragraph.add_run(text[last:match.start()])
            token = match.group()
            if match.group(1):
                hyperlink(paragraph, match.group(1), match.group(2))
            elif token.startswith('http'):
                hyperlink(paragraph, token, token)
            elif match.group(3):
                paragraph.add_run(match.group(3)).bold = True
            elif match.group(4):
                paragraph.add_run(match.group(4))
            elif token in targets:
                hyperlink(paragraph, token, '#' + targets[token])
            else:
                paragraph.add_run(token)
            last = match.end()
        paragraph.add_run(text[last:])

    # 文内目录使用实际书签，不依赖阅读器更新域，也不生成未经核验的页码。
    nav_added = False

    def navigation():
        title = doc.add_paragraph('阅读导航')
        title.runs[0].bold = True
        title.paragraph_format.keep_with_next = True
        bookmark(title, 'contents')
        headings = [line[3:] for line in lines if line.startswith('## ')]
        for start in range(0, len(headings), 2):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(4)
            for i, heading in enumerate(headings[start:start+2]):
                if i:
                    paragraph.add_run('　　')
                hyperlink(paragraph, heading, '#' + targets[heading])
        doc.add_paragraph()

    index = 0
    main_headings = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        if line.startswith('<!--'):
            index += 1
            continue
        illustration = re.fullmatch(r'!\[([^]]*)\]\(([^)]+)\)', line)
        if illustration:
            target = source.parent / illustration.group(2)
            if not target.is_file():
                target = ROOT / '决赛提交' / illustration.group(2)
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.keep_with_next = True
            paragraph.alignment = 1
            paragraph.add_run().add_picture(str(target), width=Cm(17))
            caption = doc.add_paragraph(illustration.group(1), 'Caption')
            caption.alignment = 1
            label = re.match(r'图\d+', illustration.group(1))
            if label:
                bookmark(paragraph, targets[label.group()])
            index += 1
            continue
        if line.startswith('```'):
            code = []
            index += 1
            while index < len(lines) and not lines[index].startswith('```'):
                code.append(lines[index])
                index += 1
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(.25)
            run = paragraph.add_run('\n'.join(code))
            run.font.name = 'DejaVu Sans Mono'
            run.font.size = Pt(8.5)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'F5F5F5')
            paragraph._p.get_or_add_pPr().append(shade)
            index += 1
            continue
        if line.startswith("| "):
            rows = []
            while index < len(lines) and lines[index].startswith("|"):
                value = lines[index]
                if not re.fullmatch(r"[| :\-]+", value):
                    rows.append([c.strip() for c in value.strip("|").split("|")])
                index += 1
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.style = "Table Grid"
            widths = {3: [6.2, 5.4, 5.4], 4: [5.3, 1.5, 5.1, 5.1],
                      5: [1.4, 3.9, 3.9, 3.9, 3.9],
                      6: [1.2, 3.16, 3.16, 3.16, 3.16, 3.16],
                      7: [1.2, 2.1, 2.0, 2.5, 2.1, 4.8, 2.3]}.get(len(rows[0]))
            if len(rows[0]) == 7 and rows[0][-1] == '成本 / 元':
                widths = [1.1, 2.8, 2.4, 2.8, 2.5, 2.9, 2.5]
            elif len(rows[0]) == 4 and rows[0][0] == '配置':
                widths = [3.2, 4.6, 4.6, 4.6]
            elif len(rows[0]) == 3 and rows[0][1] == '任务与实际调用标识':
                widths = [2.0, 6.2, 8.8]
            if widths:
                table.autofit = False
                for col, width in zip(table.columns, widths):
                    col.width = Cm(width)
            for ri, values in enumerate(rows):
                cells = table.add_row().cells
                for ci, (cell, value) in enumerate(zip(cells, values)):
                    if widths:
                        cell.width = Cm(widths[ci])
                    cell.text = ''
                    inline(cell.paragraphs[0], value)
                    shading = OxmlElement('w:shd')
                    shading.set(qn('w:fill'), 'EDEDED' if ri == 0 else 'FFFFFF')
                    cell._tc.get_or_add_tcPr().append(shading)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(3)
                        if ri == 0 or (len(rows) <= 6 and ri < len(rows) - 1):
                            paragraph.paragraph_format.keep_with_next = True
                        for run in paragraph.runs:
                            run.font.size = Pt(9)
                            run.bold = ri == 0
                            run.font.color.rgb = RGBColor(0, 0, 0)
                props = table.rows[-1]._tr.get_or_add_trPr()
                props.append(OxmlElement("w:cantSplit"))
                if ri == 0:
                    props.append(OxmlElement("w:tblHeader"))
            doc.add_paragraph()
            continue
        if line.startswith("# "):
            title = doc.add_paragraph(style='Title')
            label = plain(line[2:])
            if '：' in label:
                project, subtitle = label.split('：',1)
                title.add_run(project).font.size = Pt(30)
                title.add_run('\n')
                title.add_run(subtitle).font.size = Pt(21)
            else:
                title.add_run(label)
        elif line.startswith("## "):
            if not nav_added:
                navigation()
                nav_added = True
            heading = doc.add_heading(plain(line[3:]), level=1)
            # 实验单独起页；方法章节自然衔接，避免短小节独占一页。
            heading.paragraph_format.page_break_before = line.startswith('## 六、')
            main_headings += 1
            bookmark(heading, targets[line[3:]])
        elif line.startswith("### "):
            heading = doc.add_heading(plain(line[4:]), level=2)
            heading.paragraph_format.page_break_before = line.startswith('### 6.5 ')
        elif line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            inline(paragraph, line[2:])
            paragraph.paragraph_format.keep_with_next = index + 1 < len(lines) and lines[index+1].startswith('- ')
        elif re.match(r'^表\d+[　\s]', line):
            paragraph = doc.add_paragraph(plain(line), 'Caption')
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_after = Pt(4)
            if line.startswith('表2　'):
                paragraph.paragraph_format.page_break_before = True
            bookmark(paragraph, targets[re.match(r'表\d+', line).group()])
        else:
            paragraph = doc.add_paragraph()
            inline(paragraph, line)
            if main_headings:
                paragraph.paragraph_format.first_line_indent = Cm(.74)
            if line.endswith('：'):
                paragraph.paragraph_format.keep_with_next = True
        index += 1
    doc.core_properties.title = "word2jats 技术方案说明书"
    doc.core_properties.author = "JiangLab"
    doc.save(OUT / "技术方案说明书.docx")


def slides(tag):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    evidence = json.loads((ROOT / "reports/finals-closeout" / (tag + ".json")).read_text())
    rows = evidence["samples"]
    official = [r for r in rows if r["group"] != "external"]
    times = [r["wall_seconds"] for r in rows]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    dark, teal, grey = "173B3D", "14666B", "647572"

    def text(slide, value, x, y, width, height, size=23, color=dark, bold=False):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        tf = box.text_frame
        tf.word_wrap = True
        for i, line in enumerate(value.split("\n")):
            paragraph = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            paragraph.text = line
            paragraph.font.name = "Noto Sans CJK SC"
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color)
            paragraph.space_after = Pt(14)
        return box

    def page(title, subtitle):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("FAFAF6")
        strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12), prs.slide_height)
        strip.fill.solid()
        strip.fill.fore_color.rgb = RGBColor.from_string(teal)
        strip.line.fill.background()
        text(slide, title, .65, .48, 12, .8, 31, bold=True)
        text(slide, subtitle, .68, 1.35, 12, .6, 15, grey)
        text(slide, "JiangLab  /  word2jats", .68, 7.02, 8, .25, 10, grey)
        text(slide, f"{len(prs.slides):02d}", 12.05, 7.0, .6, .3, 11, grey)
        return slide

    slide = page("让 Word 稿件成为可核对的出版数据", "学术期刊结构化技术创新大赛 · 选题一 · JiangLab")
    text(slide, "word2jats", .8, 2.15, 11, 1, 56, teal, True)
    text(slide, "模型理解结构，程序取回原文，编辑确认结果。", .85, 3.65, 11.6, .8, 28)
    text(slide, "Word → JATS 1.3 XML + 原始图片 + 校样与复核记录", .85, 5.15, 11.6, .9, 23)
    slide.notes_slide.notes_text_frame.text = "开场 25 秒：我们关注的不是把文档换个后缀，而是把论文隐含的结构转为出版系统可用的数据，同时让编辑能够核对。"

    slide = page("把结构判断与内容取回分开", "源对象图记录原稿事实，模型回答用于定位与结构组织")
    labels = [("原稿事实", "文字范围\n图表与公式\n格式与编号"), ("结构理解", "文首与正文\n文献与引文\n表格列结构"),
              ("确定性生成", "按源位置取字\n组装 JATS\n输出原始媒体"), ("检查与核对", "内容覆盖与来源\nDTD 与引用\n编辑复核")]
    for i, (title, detail) in enumerate(labels):
        x = .7 + i * 3.15
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.5), Inches(2.85), Inches(2.8))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor.from_string("EDF3F0")
        box.line.fill.background()
        text(slide, title, x+.16, 2.7, 2.55, .45, 23, teal, True)
        text(slide, detail, x+.16, 3.42, 2.55, 1.8, 20)
    text(slide, "文首元信息采用受约束的局部 JATS 生成；正文与文献按源位置取回。", .85, 5.83, 12, .8, 18, grey)
    slide.notes_slide.notes_text_frame.text = "用 55 秒解释分工。重点说明模型摘抄只是定位线索，最终文字从 Word 取回。文首是明确的例外，用检查约束，不声称模型完全不生成文字。"

    slide = page("围绕出版内容做结构化处理", "不只检查 XML 能否打开，还检查内容去向与对象关系")
    text(slide, "参考文献\n两路识别条目边界，逐条并行拆字段；\n字段不完整时保留整条原文。", .8, 2.3, 5.8, 2.0, 24)
    text(slide, "公式与媒体\nOMML 转 MathML，图像按原字节交付；\n不凭图片猜出论文内容。", 7, 2.3, 5.6, 2, 24)
    text(slide, "处理效率\n引文、正文、文献边界同时推进；\n后续任务只等待真正需要的输入。", .8, 4.7, 5.8, 2, 24)
    text(slide, "内容检查\n原稿内容去了哪里，输出内容来自哪里；\n不确定项交给有原文依据的复核。", 7, 4.7, 5.6, 2, 24)
    slide.notes_slide.notes_text_frame.text = "用 50 秒说明文献、公式、并发和双向内容检查。避免逐个背模块名，把技术与真实内容对应起来。"

    slide = page("14 例完整转换验证", "10 例官方样例 + 4 例初赛评委测试稿；均已用于开发，不作为独立盲测")
    metrics = [(f"{sum(r['dtd_valid'] for r in rows)}/14", "JATS DTD 合法"),
               (f"{sum(r['delivered'] for r in official)}/10", "官方样例\n自动检查通过"),
               (str(sum(r['references'] for r in rows)), "参考文献条目\n逐例数量核对"),
               (f"{sum(r['omml_source'] for r in rows)}", "OMML 公式\n转换文本完整核对")]
    for i, (value, label) in enumerate(metrics):
        x = .8 + i * 3.13
        text(slide, value, x, 2.25, 3, 1.1, 48, teal, True)
        text(slide, label, x, 3.5, 2.8, 1.1, 19)
    text(slide, f"逐篇全新调用：{min(times):.0f}–{max(times):.0f} 秒，中位数 {statistics.median(times):.0f} 秒", .8, 5.05, 12, .8, 25)
    text(slide, "自动通过不等于排版已无需校订；复杂压平表格、旧式对象等仍进入人工核对。", .8, 6.03, 12, .65, 18, grey)
    slide.notes_slide.notes_text_frame.text = "用 40 秒报告结果。公式检查核对全部变换节点，不把文本检查说成数学语义形式证明。速度是逐篇冷跑，长稿仍可能超过五分钟。"

    slide = page("校样工作台：让剩余工作可定位、可记录", "现场演示：上传 → 成品预览 → 原稿核对 → 保存复核 → 下载")
    screenshot = ROOT / "reports/web-final/S03/工作台演示.png"
    if screenshot.is_file():
        slide.shapes.add_picture(str(screenshot), Inches(.75), Inches(2), height=Inches(4.7))
        text(slide, "对应原稿摘录\n逐项处理结论\n记录针对的 XML 版本\n修订 Word 后再转换", 8.15, 2.6, 4.5, 3.8, 25)
    else:
        text(slide, "原稿定位 → 编辑判断 → 保存复核记录 → 修订原稿后复验", .9, 3, 12, 2, 30)
    slide.notes_slide.notes_text_frame.text = "用 75 秒切到真实工作台演示。若复用响应，明确指给评委看“复用已有分析”。展示原文摘录、处理状态和记录随包导出，人工记录不覆盖自动检查。"

    slide = page("从一次转换，走向可复核的编辑流程", "已有可运行原型、源代码、14 例输入输出、运行文档与离线成果预览")
    text(slide, "减少全文重新标注，把编辑注意力集中到疑难内容。", .9, 2.5, 11.8, 1.2, 32, teal, True)
    text(slide, "已实现：原文定位、自动结构化、结果预览、复核记录与导出。\n下一步：以真实编辑处理时长检验效率，扩展复杂表格和旧式对象支持。", .9, 4.2, 11.8, 1.7, 24)
    text(slide, "word2jats.jianglab.work", .9, 6.3, 11, .5, 20, grey)
    slide.notes_slide.notes_text_frame.text = "用 25 秒收束：不是承诺替代所有编辑判断，而是自动完成可确定的工作，清楚交接剩余问题。当前未量化编辑节省百分比，不虚报。"
    prs.core_properties.author = "JiangLab"
    prs.core_properties.title = "word2jats 决赛答辩"
    prs.save(OUT / "JiangLab-决赛答辩.pptx")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--tag", default="codex-final-cold-r2")
    parser.add_argument("--output", type=Path, default=OUT)
    parser.add_argument("--document-only", action="store_true")
    parser.add_argument("--source", type=Path, help="说明书 Markdown 路径，用于工作稿版式检查")
    args = parser.parse_args()
    OUT = args.output
    OUT.mkdir(parents=True, exist_ok=True)
    document(args.source)
    if not args.document_only:
        slides(args.tag)
