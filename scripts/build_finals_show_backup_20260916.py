"""Build a fixed-frame presentation from verified high-resolution slide images.

This is a separate, explicitly non-editable showing fallback. The editable
presentation remains the primary deliverable; its notes and demo link are kept.
"""
from __future__ import annotations

import argparse
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches

from build_finals_slides_20260916 import refresh_document_properties


def build(source: Path, images: Path, output: Path) -> None:
    original = Presentation(source)
    frames = sorted(images.glob('slide-*.png'))
    assert len(frames) == len(original.slides) == 13
    prs = Presentation()
    prs.slide_width, prs.slide_height = original.slide_width, original.slide_height
    prs.core_properties.title = original.core_properties.title
    prs.core_properties.author = 'JiangLab'
    prs.core_properties.subject = '决赛答辩放映备用版（固定画面）'
    prs.core_properties.comments = ''
    for number, (frame, original_slide) in enumerate(zip(frames, original.slides), 1):
        with Image.open(frame) as image:
            assert image.width >= 3200 and image.height >= 1800
            assert abs(image.width / image.height - 16 / 9) < .001
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(frame), 0, 0, prs.slide_width, prs.slide_height)
        if original_slide.has_notes_slide:
            slide.notes_slide.notes_text_frame.text = original_slide.notes_slide.notes_text_frame.text
        if number == 12:
            # Limit the click target to the displayed URL, not the whole slide.
            link = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                         Inches(3.0), Inches(5.59),
                                         Inches(5.5), Inches(.61))
            link.name = '在线原型链接'
            link.fill.solid()
            link.fill.fore_color.rgb = RGBColor(255, 255, 255)
            alpha = OxmlElement('a:alpha')
            alpha.set('val', '0')
            link.fill._xPr.solidFill.srgbClr.append(alpha)
            link.line.fill.background()
            link.click_action.hyperlink.address = 'https://word2jats.jianglab.work'
    output.parent.mkdir(parents=True, exist_ok=True)
    refresh_document_properties(prs)
    prs.save(output)
    print(f'Fixed-frame showing backup: {len(prs.slides)} slides; {output}')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--source', type=Path, required=True)
    parser.add_argument('--images', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    build(args.source, args.images, args.output)
