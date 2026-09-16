"""Build the finals presentation using the original competition PPTX template.

All text, diagrams and tables are native PowerPoint objects. Screenshots are
the only raster content. This does not modify the submitted prototype or site.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
import json
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / 'references/学术期刊结构化技术创新大赛项目介绍.pptx'
BLACK, NAVY, GOLD = '171717', '193B67', 'F2AD00'
GREY, LIGHT, RULE = '59616B', 'F3F6FA', 'DAE0E8'
CN, EN = 'Microsoft YaHei', 'Arial'
W, H, L, CW = 13.333333, 7.5, .68, 11.98
PROJECT_TITLE = (ROOT / '决赛提交/技术方案说明书.md').read_text(encoding='utf-8').splitlines()[0].removeprefix('# ').strip()


def color(value):
    return RGBColor.from_string(value)


def font_run(run, size, bold=False, ink=BLACK, latin=EN):
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color(ink)
    props = run._r.get_or_add_rPr()
    props.set('lang', 'zh-CN')
    for name, value in [('a:ea', CN), ('a:cs', latin)]:
        el = props.find('{http://schemas.openxmlformats.org/drawingml/2006/main}' + name[2:])
        if el is None:
            el = OxmlElement(name)
            props.append(el)
        el.set('typeface', value)


def textbox(slide, text, x, y, w, h, size=21, *, bold=False,
            ink=BLACK, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
            wrap=False, line=1.18, after=0, latin=EN, name=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shape.name = name
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = valign
    for i, value in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(after)
        p.line_spacing = Pt(size * line)
        r = p.add_run()
        r.text = value
        font_run(r, size, bold, ink, latin)
    return shape


def rect(slide, x, y, w, h, fill='FFFFFF', stroke=None, width=.7):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if stroke:
        shape.line.color.rgb = color(stroke)
        shape.line.width = Pt(width)
    else:
        shape.line.fill.background()
    shape._element.spPr.append(OxmlElement('a:effectLst'))
    for style in shape._element.xpath('./p:style'):
        shape._element.remove(style)
    return shape


def line(slide, x1, y1, x2, y2, ink=RULE, width=1, arrow=False):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = color(ink)
    shape.line.width = Pt(width)
    shape._element.spPr.append(OxmlElement('a:effectLst'))
    for style in shape._element.xpath('./p:style'):
        shape._element.remove(style)
    if arrow:
        end = OxmlElement('a:tailEnd')
        end.set('type', 'triangle')
        end.set('w', 'sm')
        end.set('len', 'sm')
        shape._element.spPr.get_or_add_ln().append(end)
    return shape


def route(slide, points, ink=NAVY, width=1.3):
    for i, (a, b) in enumerate(zip(points, points[1:])):
        line(slide, *a, *b, ink=ink, width=width, arrow=i == len(points)-2)


def section(slide, title, x, y, w, body, size=21, h=2, leading=1.35):
    rect(slide, x, y+.03, .045, .29, GOLD)
    textbox(slide, title, x+.17, y-.015, w-.17, .42, 23, bold=True)
    textbox(slide, body, x+.17, y+.60, w-.17, h, size, wrap=True, line=leading)


def cell_border(cell, side, ink=RULE, width=.6):
    props = cell._tc.get_or_add_tcPr()
    ns = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
    for old in list(props.findall(ns + side)):
        props.remove(old)
    ln = OxmlElement('a:'+side)
    ln.set('w', str(int(width*12700)))
    sf = OxmlElement('a:solidFill')
    clr = OxmlElement('a:srgbClr')
    clr.set('val', ink)
    sf.append(clr)
    ln.append(sf)
    # CT_TableCellProperties requires borders before fill properties.
    order=['lnL','lnR','lnT','lnB','lnTlToBr','lnBlToTr','cell3D',
           'noFill','solidFill','gradFill','blipFill','pattFill','grpFill','extLst']
    wanted=order.index(side)
    for index,child in enumerate(props):
        name=etree.QName(child).localname
        if name in order and order.index(name)>wanted:
            props.insert(index,ln)
            break
    else:
        props.append(ln)


def table(slide, headers, data, x, y, widths, heights, size=18, *,
          aligns=None, group_rows=(), highlight_col=None, shaded_rows=(), category_col=None):
    graphic = slide.shapes.add_table(len(data)+1, len(headers), Inches(x), Inches(y),
                                    Inches(sum(widths)), Inches(sum(heights)))
    tbl = graphic.table
    for col, width in zip(tbl.columns, widths):
        col.width = Inches(width)
    for row, height in zip(tbl.rows, heights):
        row.height = Inches(height)
    for ri, values in enumerate([headers]+data):
        for ci, value in enumerate(values):
            cell = tbl.cell(ri,ci)
            cell.margin_left = cell.margin_right = Inches(.13)
            cell.margin_top = cell.margin_bottom = Inches(.035)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            bg = NAVY if ri == 0 else (LIGHT if ci == highlight_col or ci == category_col else ('F7F9FC' if ri in shaded_rows else 'FFFFFF'))
            cell.fill.fore_color.rgb = color(bg)
            tf = cell.text_frame
            tf.clear()
            tf.word_wrap = True
            tf.auto_size = MSO_AUTO_SIZE.NONE
            for pi, part in enumerate(str(value).split('\n')):
                p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
                p.space_before = p.space_after = Pt(0)
                p.line_spacing = Pt(size * 1.12)
                p.alignment = PP_ALIGN.CENTER if aligns is None else aligns[ci]
                if ri == 0:
                    p.alignment = PP_ALIGN.CENTER
                r=p.add_run(); r.text=part
                font_run(r, size+.4 if ri == 0 else size,
                         bold=ri == 0, ink='FFFFFF' if ri == 0 else BLACK)
            for side in ['lnL','lnR','lnT','lnB']:
                cell_border(cell,side, RULE, .45)
            if ri == 0:
                cell_border(cell,'lnT',NAVY,1.2)
                cell_border(cell,'lnB',NAVY,1.1)
            if ri in group_rows:
                cell_border(cell,'lnT','9EAEC2',1)
            if ri == len(data):
                cell_border(cell,'lnB',NAVY,1.1)
    return tbl


def fill_cell(cell, value, *, size=18, bold=False, ink=BLACK,
              align=PP_ALIGN.LEFT, bg=None):
    cell.text = value
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if bg:
        cell.fill.solid();cell.fill.fore_color.rgb=color(bg)
    for p in cell.text_frame.paragraphs:
        p.alignment=align;p.space_before=p.space_after=Pt(0);p.line_spacing=Pt(size*1.15)
        for r in p.runs:
            font_run(r,size,bold,ink)


def picture_fit(slide,path,x,y,w,h):
    from PIL import Image
    with Image.open(path) as im:
        iw,ih=im.size
    ratio=min(w/iw,h/ih)
    fw,fh=iw*ratio,ih*ratio
    return slide.shapes.add_picture(str(path),Inches(x+(w-fw)/2),Inches(y+(h-fh)/2),Inches(fw),Inches(fh))


def clean_template_metadata(prs):
    """Drop editor-only tags and broken image links, retaining embedded artwork.

    WPS template tags must not be shared by cloned PowerPoint slides. These tags
    and the template's `Target="NULL"` image links carry no displayed content.
    The embedded image is retained byte-for-byte for every removed image link.
    """
    pns='{http://schemas.openxmlformats.org/presentationml/2006/main}'
    rns='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
    stats={'tag_relationships_removed':0,'null_image_links_removed':0}
    for part in list(prs.part.package.iter_parts()):
        root=getattr(part,'_element',None)
        if root is None:
            continue
        for node in list(root.iter(pns+'tags')):
            node.getparent().remove(node)
        for node in list(root.iter(pns+'custDataLst')):
            if len(node)==0:
                node.getparent().remove(node)
        for rel in list(part.rels.values()):
            if rel.reltype.endswith('/tags'):
                part.rels.pop(rel.rId)
                stats['tag_relationships_removed']+=1
            elif rel.is_external and rel.reltype.endswith('/image') and rel.target_ref.upper()=='NULL':
                for node in root.iter():
                    if node.get(rns+'link')==rel.rId:
                        assert node.get(rns+'embed'), 'Cannot drop an image without embedded data'
                        del node.attrib[rns+'link']
                part.rels.pop(rel.rId)
                stats['null_image_links_removed']+=1
    return stats


def refresh_document_properties(prs):
    """Discard inherited title inventories and update actual slide/note counts."""
    ns = '{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}'
    for part in prs.part.package.iter_parts():
        if str(part.partname) != '/docProps/app.xml':
            continue
        root = etree.fromstring(part.blob)
        values = {'Slides': len(prs.slides),
                  'Notes': sum(slide.has_notes_slide for slide in prs.slides),
                  'HiddenSlides': sum(slide._element.get('show') == '0' for slide in prs.slides),
                  'Application': 'python-pptx'}
        for name, value in values.items():
            node = root.find(ns + name)
            if node is not None:
                node.text = str(value)
        for name in ['Words', 'Paragraphs', 'HeadingPairs', 'TitlesOfParts', 'AppVersion']:
            node = root.find(ns + name)
            if node is not None:
                root.remove(node)
        part._blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)


def build(asset_dir, output):
    # Preserve the actual template, including its masters, layouts and imagery.
    prs=Presentation(TEMPLATE)
    clean_template_metadata(prs)
    content_source=prs.slides[4]
    end_source=prs.slides[28]
    for sid in list(prs.slides._sldIdLst)[1:]:
        prs.part.drop_rel(sid.rId)
        prs.slides._sldIdLst.remove(sid)

    def copy_slide(source, indices=None):
        s=prs.slides.add_slide(source.slide_layout)
        for shape in list(s.shapes):
            s.shapes._spTree.remove(shape._element)
        bg=source._element.cSld.find('{http://schemas.openxmlformats.org/presentationml/2006/main}bg')
        if bg is not None:
            s._element.cSld.insert(0,deepcopy(bg))
        selected=list(source.shapes) if indices is None else [source.shapes[i] for i in indices]
        ns='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
        for shape in selected:
            element=deepcopy(shape._element)
            for node in element.iter():
                for attr,rid in list(node.attrib.items()):
                    if attr.startswith(ns):
                        rel=source.part.rels[rid]
                        node.set(attr,s.part.relate_to(rel.target_ref if rel.is_external else rel.target_part,rel.reltype,rel.is_external))
            s.shapes._spTree.insert_element_before(element,'p:extLst')
        return s

    def page(title, sub=None):
        s=copy_slide(content_source,[1,2,3,4])
        textbox(s,title,.67,.67,11.35,.83,32,bold=False,name=f'slide-{len(prs.slides):02}-title')
        if sub:
            textbox(s,sub,L,1.48,CW,.37,17.5,ink=GREY)
        textbox(s,f'{len(prs.slides):02}',12.12,7.15,.53,.22,10,ink=GREY,align=PP_ALIGN.RIGHT)
        return s

    def note(s,value):
        s.notes_slide.notes_text_frame.text=value

    # 01: the original cover, with only presentation-specific text replaced.
    s=prs.slides[0]
    for shape in list(s.shapes):
        s.shapes._spTree.remove(shape._element)
    textbox(s,'2026/09/16',9.08,.55,3.69,.5,24,ink=NAVY,align=PP_ALIGN.RIGHT)
    textbox(s,'学术期刊结构化技术创新大赛 · 选题一',.83,4.86,11.76,.34,18,ink='FFFFFF')
    cover_title=textbox(s,PROJECT_TITLE.replace('：','：\n',1),.83,5.33,11.76,1.25,30,bold=True,ink='FFFFFF',name='project-title')
    for i,p in enumerate(cover_title.text_frame.paragraphs):
        p.line_spacing=Pt(48 if i==0 else 40)
        for run in p.runs:
            font_run(run,40 if i==0 else 30,True,'FFFFFF')
    textbox(s,'JiangLab',.83,6.93,11.62,.35,18,ink='FFFFFF')
    note(s,'各位评委好，我们的作品是 Word2JATS，面向学术论文从 Word 到 JATS XML 的结构化转换。')

    # 02: the complete journal-directory screenshot and an explicit problem definition.
    s=page('需求分析与问题定义')
    section(s,'需求范围',L,1.92,5.05,'IMR Press 21 本期刊',22,h=.6)
    textbox(s,'研究论文、综述、社论、病例报告、\n读者来信、手术技术等多种文章类型',L+.17,3.19,5.0,.78,19.5,line=1.3)
    rect(s,L+.17,4.19,4.99,2.60,'FFFFFF',RULE)
    picture_fit(s,asset_dir/'imr-21-templates.png',L+.20,4.22,4.93,2.54)
    line(s,6.10,1.91,6.10,6.76,ink=RULE,width=.8)
    section(s,'问题定义',6.48,1.92,6.15,'',22,h=.01)
    textbox(s,'输入',6.66,2.52,1.0,.40,21,bold=True,ink=NAVY)
    textbox(s,'Word 文档（.docx）',7.78,2.52,4.83,.44,24)
    line(s,9.47,3.10,9.47,3.46,ink=NAVY,width=1.6,arrow=True)
    textbox(s,'输出',6.66,3.70,1.0,.40,21,bold=True,ink=NAVY)
    textbox(s,'JATS 1.3 XML\n＋ 附属资源包（ZIP）',7.78,3.65,4.83,1.02,25,line=1.35)
    line(s,6.65,5.03,12.61,5.03,ink=GOLD,width=1.3)
    textbox(s,'转换要求',6.66,5.31,5.9,.42,22,bold=True)
    textbox(s,'保留原文内容，识别章节与内容对象；\n建立作者、单位、图表与文献的关联；\n生成 XML，并配齐附属资源。',6.66,5.90,5.94,1.01,19.5,line=1.25)
    note(s,'我们从二十一本期刊及其文章类型理解需求。问题定义是：输入 Word 文档，输出符合 JATS 1.3 的 XML 和 ZIP 附属资源包。核心不只是识别内容，还包括章节结构、作者单位及图表文献之间的关联。')

    # 03: three clear columns and one scale line.
    s=page('样例清点与评价依据')
    cols=[L,4.86,9.04]
    section(s,'样例组成',cols[0],1.87,3.65,'01–05：5 例\nWord 文档＋上线 XML\n\nS01–S05：5 例\nWord 文档',21,h=3.1)
    section(s,'结构参考构建',cols[1],1.87,3.65,'以 Word 原文确定内容\n结合 JATS 规范整理结构\n标注字段、对象及关联\n\n上线 XML 辅助理解\n出版标记',20,h=3.3)
    section(s,'评价方法',cols[2],1.87,3.61,'按内容与关系匹配\n兼容等价的 JATS 表达\n\n计算精确率与召回率\n另行核对文字、原生公式\n及图片资源',20,h=3.3)
    line(s,L,5.91,L+CW,5.91,ink=RULE)
    values=[('10','例样例'),('64','位作者'),('167','个章节'),('35','幅图'),('46','张表'),('525','条参考文献')]
    for i,(value,label) in enumerate(values):
        xx=L+i*2.00
        textbox(s,value,xx,6.13,1.98,.45,29,bold=True,ink=NAVY,align=PP_ALIGN.CENTER)
        textbox(s,label,xx,6.65,1.98,.27,15.5,ink=GREY,align=PP_ALIGN.CENTER)
        if i:
            line(s,xx-.01,6.16,xx-.01,6.88,ink='E6EBF1',width=.6)
    note(s,'我们对十例样例统一清点内容。前五例有上线 XML，后五例只有 Word，因此以 Word 为内容依据，整理结构参考。评价按字段、对象和关系匹配，兼容不同但等价的 JATS 写法，并直接核对原文与资源。')

    # 04: route choice; the selected approach gets one pale column.
    s=page('技术选型：三类转换方案的比较','大语言模型（LLM）用于理解上下文、识别内容角色与关系。')
    rows=[['处理方式','正则匹配\n模板映射','程序编排，LLM 识别\n程序生成 JATS','LLM 规划步骤\n并调用工具'],
          ['主要优势','处理快\n资源消耗低','上下文识别\n与流程控制相结合','流程灵活，易拓展\n可自主调整处理步骤'],
          ['主要代价','依赖领域规则\n体例变化难维护','控制识别误差\n与调用耗时','多轮调用，耗时\n与资源更难预估'],
          ['适用特点','格式约束明确\n体例稳定','内容类别明确\n排版写法多样','处理步骤不确定\n需要探索']]
    table(s,['比较项','规则驱动转换','大模型辅助转换','智能体自主转换'],rows,L,2.04,[1.48,3.50,3.50,3.50],[.55,.83,.83,.83,.83],19.5,highlight_col=2)
    rect(s,L,6.14,.055,.56,GOLD)
    textbox(s,'本项目选择：大模型辅助转换',L+.19,6.12,11.65,.40,23,bold=True)
    textbox(s,'LLM 判断角色与关联，程序解析原稿、组织任务并生成 JATS XML。',L+.19,6.58,11.65,.29,17.5)
    note(s,'我们比较了规则驱动、大模型辅助和智能体自主处理三条路线。规则快，但需要维护体例规则；智能体灵活，但多轮调用耗时难预测。我们选择第二条路线，让大模型理解上下文，由程序控制流程与输出。')

    # 05: native editable process diagram, not a screenshot.
    s=page('总体流程：从 Word 解析到 JATS 输出')
    phases=[('Word 解析','程序','文字、样式与编号\n表格、图片与公式\n记录原稿位置'),
            ('语义识别','LLM','文首与正文层级\n图表、文献及引文\n识别角色与关系'),
            ('定位与组装','程序','定位原文范围\n组织章节与对象\n连接引用关系'),
            ('JATS 生成','程序','生成 XML 与媒体\n核对格式与引用\n检查内容及资源'),
            ('校订与下载','网页','预览与原稿对照\n修改文章出版信息\n下载转换成果')]
    for i,(title,role,body) in enumerate(phases):
        xx=L+i*2.45
        textbox(s,role,xx,1.96,2.18,.30,17,ink=NAVY,align=PP_ALIGN.CENTER)
        rect(s,xx,2.45,2.18,2.73,fill=LIGHT if i==1 else 'FFFFFF',stroke=RULE)
        rect(s,xx,2.45,2.18,.045,GOLD if i==1 else NAVY)
        textbox(s,title,xx+.13,2.82,1.92,.45,22,bold=True,align=PP_ALIGN.CENTER)
        textbox(s,body,xx+.10,3.61,1.98,1.17,17.5,align=PP_ALIGN.CENTER,line=1.48)
        if i<4:
            line(s,xx+2.21,3.65,xx+2.41,3.65,ink=NAVY,width=1.5,arrow=True)
    route(s,[(1.77,5.24),(1.77,5.70),(6.67,5.70),(6.67,5.24)],ink=NAVY)
    textbox(s,'原文与资源直接进入组装',2.56,5.85,5.24,.37,19,ink=NAVY,align=PP_ALIGN.CENTER)
    textbox(s,'文首题名、作者及机构信息：限定原稿区域生成局部 JATS，并作专门校验。',L,6.48,CW,.36,18.5)
    note(s,'总体流程先解析 Word，保留原稿位置；再由大模型判断结构。程序根据识别结果取回原文，组装 JATS 并检查，最后进入网页校订。文首元信息采用限定区域的局部 JATS 生成与专门校验。')

    # 06: one small example explains the central method.
    s=page('核心方法：LLM 识别结构，程序取回原文')
    parts=[(L,3.18,'原文'),(4.10,3.48,'LLM 识别'),(7.82,4.84,'程序组装')]
    for xx,ww,title in parts:
        rect(s,xx,2.02,ww,3.30,'FFFFFF',RULE)
        rect(s,xx,2.02,ww,.51,NAVY)
        textbox(s,title,xx+.16,2.08,ww-.32,.32,21,bold=True,ink='FFFFFF',align=PP_ALIGN.CENTER)
    line(s,3.89,3.67,4.05,3.67,ink=NAVY,width=1.3,arrow=True)
    line(s,7.61,3.67,7.78,3.67,ink=NAVY,width=1.3,arrow=True)
    textbox(s,'2. Materials\nand Methods',L+.22,3.08,2.74,1.15,25,bold=True,ink=NAVY,line=1.25)
    textbox(s,'角色：章节标题\n位置：doc/p30\n原文片段：',4.31,2.94,3.08,1.46,20,line=1.45)
    textbox(s,'2. Materials and Methods',4.31,4.52,3.11,.37,17.5,ink=NAVY)
    code='<sec id="S2">\n  <title>\n    2. Materials and Methods\n  </title>\n  <p>……</p>\n</sec>'
    textbox(s,code,8.02,2.85,4.45,2.08,18,ink=NAVY,latin='Consolas',line=1.28)
    textbox(s,'精确匹配   →   表示差异归一化   →   上下文与顺序消歧',L,5.81,CW,.45,23,bold=False,align=PP_ALIGN.CENTER)
    textbox(s,'依据原稿位置取回标题、正文与文献字段，保留文字、行内格式及来源位置。',L,6.43,CW,.43,19.5)
    note(s,'以章节标题为例，大模型返回角色、位置与原文片段，程序再定位到实际原稿，取回文字并组装 XML。重复片段结合上下文和顺序消歧。同一位置记录继续用于内容检查与原稿对照。')

    # 07: complete methods, readable as one native table.
    s=page('图表、公式与参考文献的转换')
    rows=[['原生表格','读取行列与合并关系','JATS 表格网格'],
          ['文本排表','将原文片段分配到对应单元格','JATS 表格网格'],
          ['图片与复合图','保留资源，关联图号、图注与子图','图对象与配套图片'],
          ['数学公式','原生公式转 MathML；图像或复杂排版\n公式保留可读载体','行内、行间公式及编号'],
          ['参考文献','识别条目边界，再提取著录字段','结构化或混合著录'],
          ['正文引用','按编号或作者—年份匹配目标','连接文献、图表的引用']]
    table(s,['内容','处理方法','输出'],rows,L,1.91,[2.12,6.22,3.64],[.55,.62,.62,.68,.85,.68,.68],19.5,
          aligns=[PP_ALIGN.CENTER,PP_ALIGN.LEFT,PP_ALIGN.LEFT],category_col=0)
    note(s,'不同载体分别处理：原生表格读取网格，文本排表恢复行列；原生公式转为 MathML，图像和复杂排版保留可读载体。参考文献先确定条目，再提取字段，正文引用与目标一起建立。')

    # 08: dependency graph; all edges represent actual prerequisites.
    s=page('文档切分与并行调度')
    def node(xx,yy,ww,hh,title,body=None):
        rect(s,xx,yy,ww,hh,LIGHT, RULE)
        if body:
            textbox(s,title,xx+.10,yy+.10,ww-.20,.3,18,bold=True,align=PP_ALIGN.CENTER)
            textbox(s,body,xx+.10,yy+.46,ww-.20,.3,16,align=PP_ALIGN.CENTER)
        else:
            textbox(s,title,xx+.08,yy+.08,ww-.16,hh-.16,18,bold=False,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
    # Draw connectors before nodes.
    line(s,1.94,3.47,2.14,3.47,ink=NAVY,width=1.25)
    line(s,2.14,2.24,2.14,5.77,ink=NAVY,width=1.25)
    for yy in [2.24,3.42,4.59,5.77]:
        line(s,2.14,yy,2.39,yy,ink=NAVY,width=1.25,arrow=True)
    route(s,[(4.43,2.24),(4.66,2.24),(4.66,2.76),(4.94,2.76)])
    route(s,[(4.43,3.42),(4.73,3.42),(4.73,2.98),(4.94,2.98)])
    route(s,[(4.43,4.59),(4.83,4.59),(4.83,3.20),(4.94,3.20)])
    route(s,[(4.43,4.59),(4.65,4.59),(4.65,4.92),(4.94,4.92)])
    route(s,[(6.78,2.98),(7.18,2.98)])
    route(s,[(8.02,3.42),(8.02,5.23)])
    route(s,[(6.78,4.92),(6.98,4.92),(6.98,5.68),(7.18,5.68)])
    route(s,[(4.43,5.77),(6.65,5.77),(6.65,5.91),(7.18,5.91)])
    node(L,3.04,1.26,.86,'原稿\n记录')
    for yy,head,body in [(1.82,'文首识别','元信息与摘要'),(3.0,'正文识别','角色、层级与图表'),(4.17,'文献边界','条目定位与归并'),(5.35,'正文引文','位置与目标线索')]:
        node(2.39,yy,2.04,.84,head,body)
    node(4.94,2.56,1.84,.84,'内容角色\n归并')
    node(7.18,2.56,1.68,.84,'文本表格\n恢复')
    node(4.94,4.50,1.84,.84,'文献字段','各条并行提取')
    node(7.18,5.23,1.68,.92,'最终组装\n与检查')
    line(s,9.27,1.88,9.27,6.45,ink=RULE)
    section(s,'长文处理',9.57,2.0,3.05,'按输入预算分窗\n保留相邻上下文\n按原稿位置归并',19,h=1.55)
    section(s,'并发控制',9.57,4.38,3.05,'LLM 并发上限：32\n异常请求限次重试',19,h=1.14)
    note(s,'耗时优化主要依靠任务切分和依赖调度。文首、正文、文献边界和引文并行启动；文献边界一旦确认，就开始逐条字段提取，无需等待正文结束。长文分窗处理，所有 LLM 请求共享并发上限。')

    # 09: verified values, grouped without a fabricated overall score.
    s=page('转换质量：内容、结构与关系','10 例样例 · 两种大模型配置的分项召回率与完整性检查')
    rows=[['题名、作者及单位关联、通讯信息、ORCID、日期','各项 100%','各项 100%'],
          ['摘要、关键词、章节及顺序','各项 100%','各项 100%'],
          ['术语释义与文末声明','各项 100%','各项 100%'],
          ['图对象、图注、图片归属与图表引用','各项 100%','各项 100%'],
          ['表对象、表题 / 表注','100% / 93.94%','100% / 93.94%'],
          ['表格网格位置与内容','77.34%','93.26%'],
          ['原生公式数学树 / 公式图片','26/26；2/2','26/26；2/2'],
          ['文献条目 / 文献著录字段','100% / 99.06%','100% / 96.83%'],
          ['正文—文献引用关系','92.41%','94.12%'],
          ['全文词项保留率','99.73%','99.71%'],
          ['JATS DTD 合法文档','10/10','10/10'],
          ['导出媒体原字节一致性','50/50','50/50']]
    table(s,['评价项','Qwen','DeepSeek'],rows,L,1.96,[7.16,2.41,2.41],[.46]+[.374]*12,18,
          aligns=[PP_ALIGN.LEFT,PP_ALIGN.CENTER,PP_ALIGN.CENTER],group_rows=[4,8,10],shaded_rows=[4,5,6,7,10,11,12])
    note(s,'质量评价覆盖文首、正文、图表、公式、参考文献和关联关系。两种配置在主要对象上表现接近，表格网格与文献字段存在差异：本次 DeepSeek 的表格网格召回更高，Qwen 的文献字段召回更高。整体文字保留率均超过百分之九十九点七。')

    # 10: one four-column table, category spans and meaningful rules.
    s=page('转换效率与调用成本','10 例样例 · Qwen 与 DeepSeek 两种大模型配置')
    rows=[['大模型配置','主流程','qwen3.7-plus','deepseek-flash'],
          ['','文首元信息','qwen3.8-max','deepseek-flash'],
          ['转换耗时\n秒','平均值','135.90','58.27'],
          ['','中位数','92.31','27.34'],
          ['Token 用量\n10 例合计','输入合计','5,815,531','5,302,342'],
          ['','  其中：缓存命中','2,660,992','5,175,210'],
          ['','  其中：缓存未命中','3,154,539','127,132'],
          ['','输出','928,179','993,918'],
          ['','输入＋输出','6,743,710','6,296,260'],
          ['调用费用\n元，按 Token 折算','篇均','1.1839','0.8413'],
          ['','10 例合计','11.8391','8.4126']]
    tbl=table(s,['类别','指标','Qwen','DeepSeek'],rows,L,1.96,[1.96,3.02,3.50,3.50],[.46]+[.401]*11,18,
              aligns=[PP_ALIGN.CENTER,PP_ALIGN.LEFT,PP_ALIGN.CENTER,PP_ALIGN.CENTER],group_rows=[3,5,10])
    for start,end,label in [(1,2,'大模型配置'),(3,4,'转换耗时\n秒'),(5,9,'Token 用量\n10 例合计'),(10,11,'调用费用\n元（折算）')]:
        first=tbl.cell(start,0);first.merge(tbl.cell(end,0))
        fill_cell(first,label,size=17,bold=True,bg=LIGHT,align=PP_ALIGN.CENTER)
    note(s,'同一流程下，我们比较两种大模型配置。DeepSeek 平均约五十八秒，Qwen 约一百三十六秒；篇均调用费用分别约零点八四元和一点一八元。输入、输出与缓存用量均来自接口返回，按统一说明的计价口径折算。')

    # 11: actual saved task, no simulated application UI.
    s=page('原稿对照与信息校订')
    screenshot=asset_dir/'product-publication.png'
    rect(s,L,1.83,8.46,4.99,'FFFFFF',RULE)
    picture_fit(s,screenshot,L+.035,1.865,8.39,4.92)
    for yy,title,body in [(1.97,'预览与定位','按章节、图表、公式\n和文献浏览结果，\n查看引用目标。'),
                          (3.63,'对照与修改','对照 Word 原稿，\n修改题名、作者与机构，\n补充期刊和 DOI。'),
                          (5.27,'保存与交付','修改写入 XML；\n重新检查，预览与\n下载同步更新。')]:
        section(s,title,9.56,yy,3.04,body,18.5,h=1.10,leading=1.22)
    note(s,'网页支持内容预览、原稿对照和信息校订。题名、作者及机构信息可以修改，期刊和 DOI 可以补充。保存后更新的是实际 XML，预览和下载同步更新，让人工校订直接作用于交付成果。')

    # 12: a quiet, clickable handoff to the live demonstration.
    s=page('在线演示')
    textbox(s,'Word2JATS',L,1.87,CW,.78,42,bold=True,ink=NAVY)
    steps=['上传 Word','转换与预览','对照与校订','下载成果']
    for i,label in enumerate(steps):
        xx=L+i*3.10
        textbox(s,f'0{i+1}',xx,3.19,.55,.44,21,bold=True,ink=NAVY)
        textbox(s,label,xx+.63,3.15,2.15,.46,23,bold=True)
        if i<3:
            line(s,xx+2.66,3.42,xx+2.97,3.42,ink=NAVY,width=1.25,arrow=True)
    line(s,L,4.12,L+CW,4.12,ink=GOLD,width=1.25)
    textbox(s,'交付文件',L,4.55,2.1,.42,22,bold=True)
    textbox(s,'JATS XML ＋ 附属资源包（ZIP）',3.0,4.55,9.65,.47,25)
    textbox(s,'在线原型',L,5.64,2.1,.42,22,bold=True)
    link=textbox(s,'word2jats.jianglab.work',3.0,5.59,9.65,.61,30,ink=NAVY)
    link_run=link.text_frame.paragraphs[0].runs[0]
    link_run.hyperlink.address='https://word2jats.jianglab.work'
    for name in ['a:hlink', 'a:folHlink']:
        for theme in (part for part in prs.part.package.iter_parts() if str(part.partname).startswith('/ppt/theme/')):
            root=etree.fromstring(theme.blob)
            ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
            for color_node in root.findall('.//'+name,ns):
                for child in list(color_node): color_node.remove(child)
                rgb=OxmlElement('a:srgbClr');rgb.set('val',NAVY);color_node.append(rgb)
            theme._blob=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
    note(s,'下面通过在线原型展示从上传 Word，到查看、校订，再到下载 XML 和配套图片的完整过程。')

    # 13: reuse the original template's closing slide without redesign.
    s=copy_slide(end_source)
    note(s,'谢谢各位评委。欢迎提问。')

    prs.core_properties.title=PROJECT_TITLE
    prs.core_properties.subject='Word 稿件的 JATS 结构化转换'
    prs.core_properties.author='JiangLab'
    prs.core_properties.keywords='Word2JATS, JATS, JiangLab'
    prs.core_properties.comments=''
    clean_template_metadata(prs)
    refresh_document_properties(prs)
    # Guard the two delivery regressions: authoritative title and OOXML order.
    assert PROJECT_TITLE in [sh.text.replace('\n','') for sh in prs.slides[0].shapes if sh.has_text_frame]
    tc_order=['lnL','lnR','lnT','lnB','lnTlToBr','lnBlToTr','cell3D',
              'noFill','solidFill','gradFill','blipFill','pattFill','grpFill','extLst']
    for sl in prs.slides:
        for tcpr in sl._element.xpath('.//a:tcPr'):
            positions=[tc_order.index(etree.QName(child).localname) for child in tcpr]
            assert positions == sorted(positions), 'Invalid table-cell property order'
    output.parent.mkdir(parents=True,exist_ok=True)
    prs.save(output)
    checks=[]
    for si,sl in enumerate(prs.slides,1):
        for shape in sl.shapes:
            if shape.left < -1 or shape.top < -1 or shape.left+shape.width > prs.slide_width+200 or shape.top+shape.height > prs.slide_height+200:
                checks.append({'slide':si,'name':shape.name,'error':'outside_slide'})
    report={'slides':len(prs.slides),'out_of_bounds':checks,
            'editable_tables':sum(sh.has_table for sl in prs.slides for sh in sl.shapes),
            'screenshots':sum(sh.shape_type==13 for sl in prs.slides for sh in sl.shapes)}
    print(json.dumps(report,ensure_ascii=False))
    if checks:
        raise SystemExit('Objects outside the slide canvas')


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--assets',type=Path,default=ROOT/'决赛提交/assets/答辩')
    parser.add_argument('--output',type=Path,default=ROOT/'决赛提交/JiangLab-决赛答辩.pptx')
    args=parser.parse_args()
    build(args.assets.resolve(),args.output.resolve())
