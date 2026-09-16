"""Bounded artifact regressions, supplementing Open XML SDK and visual checks."""
from __future__ import annotations

import argparse
from collections import Counter
import json
from pathlib import Path, PurePosixPath
import posixpath
import re
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation

NS = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
}
RNS = '{' + NS['r'] + '}'


def normalized(text):
    return re.sub(r'\s+', '', text)


def slide_text(slide):
    return normalized(''.join(slide._element.xpath('.//a:t/text()')))


def check(path: Path, title: str, reference: Path | None, fixed: bool):
    prs = Presentation(path)
    assert len(prs.slides) == 13
    assert prs.core_properties.title == title
    if not fixed:
        assert normalized(title) in slide_text(prs.slides[0])
        assert 'THEEND谢谢' in slide_text(prs.slides[-1])
    if reference:
        old = Presentation(reference)
        assert len(old.slides) == len(prs.slides)
        for index, (before, after) in enumerate(zip(old.slides, prs.slides), 1):
            assert slide_text(before) == slide_text(after), f'Content changed on slide {index}'
            assert before.notes_slide.notes_text_frame.text == after.notes_slide.notes_text_frame.text
    for index, slide in enumerate(prs.slides, 1):
        ids = slide._element.xpath('.//p:cNvPr/@id')
        assert len(ids) == len(set(ids)), f'Duplicate shape IDs on slide {index}'
        assert len(slide.shapes) > 0
        if fixed:
            pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
            assert len(pictures) == 1
            image = pictures[0]
            assert (image.left, image.top, image.width, image.height) == (0, 0, prs.slide_width, prs.slide_height)
    external = []
    checked_relations = 0
    with ZipFile(path) as archive:
        names = archive.namelist()
        assert len(names) == len(set(names)), 'Duplicate ZIP part names'
        assert archive.testzip() is None
        assert not any(name.startswith('ppt/tags/') for name in names)
        props = etree.fromstring(archive.read('docProps/app.xml'))
        props_ns = '{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}'
        assert props.findtext(props_ns + 'Slides') == '13'
        assert props.findtext(props_ns + 'Notes') == '13'
        all_names = set(names)
        relation_maps = {}
        for name in names:
            if not name.endswith('.rels'):
                continue
            root = etree.fromstring(archive.read(name))
            if name == '_rels/.rels':
                source = ''
            else:
                parent = PurePosixPath(name).parent.parent
                source = str(parent / PurePosixPath(name).name.removesuffix('.rels'))
            rels = {}
            for rel in root:
                rid, target = rel.get('Id'), rel.get('Target')
                assert rid not in rels, f'Duplicate relationship in {name}'
                rels[rid] = rel
                assert target and target.upper() != 'NULL', f'Invalid target in {name}'
                if rel.get('TargetMode') == 'External':
                    external.append(target)
                else:
                    resolved = posixpath.normpath(posixpath.join(posixpath.dirname(source), target)).lstrip('/')
                    assert resolved in all_names, f'Missing part: {source} -> {target}'
                checked_relations += 1
            relation_maps[source] = rels
        for name in names:
            if not name.endswith('.xml'):
                continue
            root = etree.fromstring(archive.read(name))
            for element in root.iter():
                for attribute, value in element.attrib.items():
                    if attribute.startswith(RNS):
                        assert value in relation_maps.get(name, {}), f'Missing relationship {name}: {value}'
        assert Counter(external) == Counter(['https://word2jats.jianglab.work']), external
    return {'file': str(path), 'slides': 13, 'relationships': checked_relations,
            'title_matches_proposal': True,
            'text_comparison': 'unchanged' if reference else 'not_applicable_fixed_frame' if fixed else 'not_requested',
            'external_links': external, 'package_integrity': 'passed'}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('presentation', type=Path)
    parser.add_argument('--proposal', type=Path, default=Path('决赛提交/技术方案说明书.md'))
    parser.add_argument('--reference', type=Path)
    parser.add_argument('--fixed', action='store_true')
    args = parser.parse_args()
    title = args.proposal.read_text().splitlines()[0].removeprefix('# ').strip()
    print(json.dumps(check(args.presentation, title, args.reference, args.fixed), ensure_ascii=False, indent=2))
