"""由技术说明 Markdown 生成 Word，并用同批验证结果生成答辩 PPT。"""
from pathlib import Path
import argparse
import json
import re
import statistics

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "决赛提交"


def plain(text):
    text = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1（\2）", text)
    return text.replace("**", "").replace("`", "")


def document():
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(2.3)
    section.footer_distance = Cm(.8)
    section.left_margin, section.right_margin = Cm(2.0), Cm(2.0)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Noto Sans CJK SC"
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.keep_together = True
    for name, size in (("Title", 25), ("Heading 1", 17), ("Heading 2", 12.5)):
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor.from_string("164D50")
    section.header.paragraphs[0].text = "JiangLab   /   word2jats                                      决赛技术方案说明书"
    section.header.paragraphs[0].style = "Caption"
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("JiangLab  ·  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    lines = (ROOT / "决赛提交/技术方案说明书.md").read_text().splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        if line.startswith("| "):
            rows = []
            while index < len(lines) and lines[index].startswith("|"):
                value = lines[index]
                if not re.fullmatch(r"[| :\-]+", value):
                    rows.append([plain(c.strip()) for c in value.strip("|").split("|")])
                index += 1
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.style = "Light Shading Accent 1"
            for ri, values in enumerate(rows):
                cells = table.add_row().cells
                for cell, value in zip(cells, values):
                    cell.text = value
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(3)
                        if len(rows) <= 6 and ri < len(rows) - 1:
                            paragraph.paragraph_format.keep_with_next = True
                        for run in paragraph.runs:
                            run.font.size = Pt(9)
                            run.bold = ri == 0
                props = table.rows[-1]._tr.get_or_add_trPr()
                props.append(OxmlElement("w:cantSplit"))
                if ri == 0:
                    props.append(OxmlElement("w:tblHeader"))
            doc.add_paragraph()
            continue
        if line.startswith("# "):
            doc.add_paragraph(plain(line[2:]), "Title")
        elif line.startswith("## "):
            doc.add_heading(plain(line[3:]), level=1)
        elif line.startswith("### "):
            doc.add_heading(plain(line[4:]), level=2)
        elif line.startswith("- "):
            doc.add_paragraph(plain(line[2:]), "List Bullet")
        else:
            doc.add_paragraph(plain(line))
        index += 1
    doc.core_properties.title = "word2jats 技术方案说明书"
    doc.core_properties.author = "JiangLab"
    doc.save(OUT / "技术方案说明书.docx")


def slides(tag):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    evidence = json.loads((ROOT / "reports/finals-closeout" / (tag + ".json")).read_text())
    rows = evidence["samples"]
    official = [r for r in rows if r["group"] != "external"]
    times = [r["wall_seconds"] for r in rows]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    dark, teal, grey = "173B3D", "14666B", "647572"

    def text(slide, value, x, y, width, height, size=23, color=dark, bold=False):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        tf = box.text_frame
        tf.word_wrap = True
        for i, line in enumerate(value.split("\n")):
            paragraph = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            paragraph.text = line
            paragraph.font.name = "Noto Sans CJK SC"
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color)
            paragraph.space_after = Pt(14)
        return box

    def page(title, subtitle):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("FAFAF6")
        strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12), prs.slide_height)
        strip.fill.solid()
        strip.fill.fore_color.rgb = RGBColor.from_string(teal)
        strip.line.fill.background()
        text(slide, title, .65, .48, 12, .8, 31, bold=True)
        text(slide, subtitle, .68, 1.35, 12, .6, 15, grey)
        text(slide, "JiangLab  /  word2jats", .68, 7.02, 8, .25, 10, grey)
        text(slide, f"{len(prs.slides):02d}", 12.05, 7.0, .6, .3, 11, grey)
        return slide

    slide = page("让 Word 稿件成为可核对的出版数据", "学术期刊结构化技术创新大赛 · 选题一 · JiangLab")
    text(slide, "word2jats", .8, 2.15, 11, 1, 56, teal, True)
    text(slide, "模型理解结构，程序取回原文，编辑确认结果。", .85, 3.65, 11.6, .8, 28)
    text(slide, "Word → JATS 1.3 XML + 原始图片 + 校样与复核记录", .85, 5.15, 11.6, .9, 23)
    slide.notes_slide.notes_text_frame.text = "开场 25 秒：我们关注的不是把文档换个后缀，而是把论文隐含的结构转为出版系统可用的数据，同时让编辑能够核对。"

    slide = page("把结构判断与内容取回分开", "源对象图记录原稿事实，模型回答用于定位与结构组织")
    labels = [("原稿事实", "文字范围\n图表与公式\n格式与编号"), ("结构理解", "文首与正文\n文献与引文\n表格列结构"),
              ("确定性生成", "按源位置取字\n组装 JATS\n输出原始媒体"), ("检查与核对", "内容覆盖与来源\nDTD 与引用\n编辑复核")]
    for i, (title, detail) in enumerate(labels):
        x = .7 + i * 3.15
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.5), Inches(2.85), Inches(2.8))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor.from_string("EDF3F0")
        box.line.fill.background()
        text(slide, title, x+.16, 2.7, 2.55, .45, 23, teal, True)
        text(slide, detail, x+.16, 3.42, 2.55, 1.8, 20)
    text(slide, "文首元信息采用受约束的局部 JATS 生成；正文与文献按源位置取回。", .85, 5.83, 12, .8, 18, grey)
    slide.notes_slide.notes_text_frame.text = "用 55 秒解释分工。重点说明模型摘抄只是定位线索，最终文字从 Word 取回。文首是明确的例外，用检查约束，不声称模型完全不生成文字。"

    slide = page("围绕出版内容做结构化处理", "不只检查 XML 能否打开，还检查内容去向与对象关系")
    text(slide, "参考文献\n两路识别条目边界，逐条并行拆字段；\n字段不完整时保留整条原文。", .8, 2.3, 5.8, 2.0, 24)
    text(slide, "公式与媒体\nOMML 转 MathML，图像按原字节交付；\n不凭图片猜出论文内容。", 7, 2.3, 5.6, 2, 24)
    text(slide, "处理效率\n引文、正文、文献边界同时推进；\n后续任务只等待真正需要的输入。", .8, 4.7, 5.8, 2, 24)
    text(slide, "内容检查\n原稿内容去了哪里，输出内容来自哪里；\n不确定项交给有原文依据的复核。", 7, 4.7, 5.6, 2, 24)
    slide.notes_slide.notes_text_frame.text = "用 50 秒说明文献、公式、并发和双向内容检查。避免逐个背模块名，把技术与真实内容对应起来。"

    slide = page("14 例完整转换验证", "10 例官方样例 + 4 例初赛评委测试稿；均已用于开发，不作为独立盲测")
    metrics = [(f"{sum(r['dtd_valid'] for r in rows)}/14", "JATS DTD 合法"),
               (f"{sum(r['delivered'] for r in official)}/10", "官方样例\n自动检查通过"),
               (str(sum(r['references'] for r in rows)), "参考文献条目\n逐例数量核对"),
               (f"{sum(r['omml_source'] for r in rows)}", "OMML 公式\n转换文本完整核对")]
    for i, (value, label) in enumerate(metrics):
        x = .8 + i * 3.13
        text(slide, value, x, 2.25, 3, 1.1, 48, teal, True)
        text(slide, label, x, 3.5, 2.8, 1.1, 19)
    text(slide, f"逐篇全新调用：{min(times):.0f}–{max(times):.0f} 秒，中位数 {statistics.median(times):.0f} 秒", .8, 5.05, 12, .8, 25)
    text(slide, "自动通过不等于排版已无需校订；复杂压平表格、旧式对象等仍进入人工核对。", .8, 6.03, 12, .65, 18, grey)
    slide.notes_slide.notes_text_frame.text = "用 40 秒报告结果。公式检查核对全部变换节点，不把文本检查说成数学语义形式证明。速度是逐篇冷跑，长稿仍可能超过五分钟。"

    slide = page("校样工作台：让剩余工作可定位、可记录", "现场演示：上传 → 成品预览 → 原稿核对 → 保存复核 → 下载")
    screenshot = ROOT / "reports/web-final/S03/工作台演示.png"
    if screenshot.is_file():
        slide.shapes.add_picture(str(screenshot), Inches(.75), Inches(2), height=Inches(4.7))
        text(slide, "对应原稿摘录\n逐项处理结论\n记录针对的 XML 版本\n修订 Word 后再转换", 8.15, 2.6, 4.5, 3.8, 25)
    else:
        text(slide, "原稿定位 → 编辑判断 → 保存复核记录 → 修订原稿后复验", .9, 3, 12, 2, 30)
    slide.notes_slide.notes_text_frame.text = "用 75 秒切到真实工作台演示。若复用响应，明确指给评委看“复用已有分析”。展示原文摘录、处理状态和记录随包导出，人工记录不覆盖自动检查。"

    slide = page("从一次转换，走向可复核的编辑流程", "已有可运行原型、源代码、14 例输入输出、运行文档与离线成果预览")
    text(slide, "减少全文重新标注，把编辑注意力集中到疑难内容。", .9, 2.5, 11.8, 1.2, 32, teal, True)
    text(slide, "已实现：原文定位、自动结构化、结果预览、复核记录与导出。\n下一步：以真实编辑处理时长检验效率，扩展复杂表格和旧式对象支持。", .9, 4.2, 11.8, 1.7, 24)
    text(slide, "word2jats.jianglab.work", .9, 6.3, 11, .5, 20, grey)
    slide.notes_slide.notes_text_frame.text = "用 25 秒收束：不是承诺替代所有编辑判断，而是自动完成可确定的工作，清楚交接剩余问题。当前未量化编辑节省百分比，不虚报。"
    prs.core_properties.author = "JiangLab"
    prs.core_properties.title = "word2jats 决赛答辩"
    prs.save(OUT / "JiangLab-决赛答辩.pptx")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--tag", default="codex-final-cold-r2")
    parser.add_argument("--output", type=Path, default=OUT)
    parser.add_argument("--document-only", action="store_true")
    args = parser.parse_args()
    OUT = args.output
    OUT.mkdir(parents=True, exist_ok=True)
    document()
    if not args.document_only:
        slides(args.tag)
