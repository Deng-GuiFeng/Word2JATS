"""验证文档导航是实际链接，字体及正文颜色可随 Word/PDF 导出。"""
import zipfile

from lxml import etree

from scripts import render_finals_documents as render


def test_document_links_and_typography(tmp_path, monkeypatch):
    source = tmp_path / 'proposal.md'
    source.write_text('# 方案\n\n在线原型：[打开](https://example.org)\n\n'
                      '## 一、架构\n\n参见表1。\n\n表1　数据\n\n'
                      '| 配置 | 结果 |\n|---|---|\n| A | 1 |\n\n'
                      '## 二、结果\n\n文字。\n', encoding='utf-8')
    monkeypatch.setattr(render, 'OUT', tmp_path)
    render.document(source)
    ns = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
          'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    with zipfile.ZipFile(tmp_path / '技术方案说明书.docx') as archive:
        body = etree.fromstring(archive.read('word/document.xml'))
        footer = etree.fromstring(archive.read('word/footer1.xml'))
        styles = etree.fromstring(archive.read('word/styles.xml'))
        relations = archive.read('word/_rels/document.xml.rels')
    bookmarks = body.xpath('//w:bookmarkStart/@w:name', namespaces=ns)
    anchors = body.xpath('//w:hyperlink/@w:anchor', namespaces=ns) + footer.xpath('//w:hyperlink/@w:anchor', namespaces=ns)
    assert len(bookmarks) == len(set(bookmarks))
    assert set(anchors) <= set(bookmarks)
    assert 'table_1' in anchors and 'contents' in bookmarks
    assert len(body.xpath('//w:hyperlink[@r:id]', namespaces=ns)) == 1
    assert b'Target="https://example.org"' in relations
    normal = styles.xpath('//w:style[@w:styleId="Normal"]/w:rPr', namespaces=ns)[0]
    assert normal.xpath('w:rFonts/@w:ascii', namespaces=ns) == ['Times New Roman']
    assert normal.xpath('w:rFonts/@w:eastAsia', namespaces=ns) == ['Noto Serif CJK SC']
    assert normal.xpath('w:color/@w:val', namespaces=ns) == ['000000']


def test_two_level_contents_body_alignment_and_table_bounds(tmp_path, monkeypatch):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    source = tmp_path/'proposal.md'
    source.write_text('# Word2JATS：技术说明\n\n## 一、目标\n\n### 1.1 方法\n\n正文内容。\n\n'
                      '| 编号 | 输入 | 输出 | 缓存命中 | 缓存未命中 | 总量 | 成本 / 元 |\n'
                      '|---|---|---|---|---|---|---|\n| 01 | 100 | 5 | 90 | 10 | 105 | 0.0001 |\n\n'
                      '## 二、验证\n\n### 2.1 结果\n\n文字。', encoding='utf-8')
    monkeypatch.setattr(render, 'OUT', tmp_path)
    render.document(source, {'section_0':2,'section_1':2,'section_2':3,'section_3':3})
    doc = Document(tmp_path/'技术方案说明书.docx')
    assert doc.styles['Normal'].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    ns = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    with zipfile.ZipFile(tmp_path/'技术方案说明书.docx') as archive:
        body = etree.fromstring(archive.read('word/document.xml'))
    assert not body.xpath('//w:t[text()="阅读导航"]', namespaces=ns)
    assert len(body.xpath('//w:hyperlink[starts-with(@w:anchor,"section_")]', namespaces=ns)) == 4
    assert all(link.xpath('.//w:t/text()',namespaces=ns) for link in body.xpath('//w:hyperlink',namespaces=ns))
    text_width = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    assert sum(col.width for col in doc.tables[0].columns) <= text_width
    assert doc.tables[0].cell(1,1).paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert doc.tables[0].cell(1,0).paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert len(body.xpath('//w:pPr/w:pageBreakBefore[not(@w:val="0")]',namespaces=ns)) == 1
